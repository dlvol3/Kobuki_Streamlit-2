"""Read, never modify, Group 4's worksheets; publish bounded model data."""
from pathlib import Path
from hashlib import sha256
import csv
import json
import sys

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
SOURCE = Path(sys.argv[1]) if len(sys.argv)>1 else ROOT.parent / "Kobuki_Group_4.xlsx"
wb = load_workbook(SOURCE, read_only=True, data_only=True)
ws = wb["Question 1"]
mapping = {"price":"C7","unit_variable_cost":"C8","annual_fixed_cost":"C9", "capex":"C15", "life":"C16", "salvage":"C17", "initial_nwc":"C23", "nwc_ratio":"C24", "tax_rate":"C26", "risk_free":"F30", "credit_spread":"F33", "beta":"F31", "market_risk_premium":"F32", "debt_weight":"C35", "equity_weight":"C34"}
inputs = {k:ws[cell].value for k,cell in mapping.items()}
inputs["base_units"] = [ws[f"{col}43"].value for col in "DEFGHIJ"]
if any(v is None for v in inputs.values()):
    raise RuntimeError("Workbook cached values are incomplete. Recalculate and save it in Excel first.")
source = {"file":SOURCE.name,"sheet":"Question 1","sha256":sha256(SOURCE.read_bytes()).hexdigest(),"input_cells":mapping,"sales_cells":"D43:J43","npv_cell":"C59","irr_cell":"C61","fcf_cells":"C56:J56"}
q2 = wb['Question 2']
presets = {
    'q1': {'label':'Q1 baseline','sales_coefficient':1.0,'fixed_cost_change':0.0,'sheet':'Question 1'},
    'q2': {'label':'Q2 baseline','sales_coefficient':1+q2['C8'].value,'fixed_cost_change':q2['C9'].value+q2['C10'].value,'sheet':'Question 2','source_cells':{'sales_uplift':'C8','project_fixed_cost_change':'C9','corporate_fixed_cost_change':'C10'}}
}
data = ROOT / "data"
data.mkdir(exist_ok=True)
(data/"case_inputs.json").write_text(json.dumps({"version":"2026-10-19-v3", "source":source,"inputs":inputs,'presets':presets,"assumptions":["Annual year-end operating cash flows.","NWC equals 0.10 of current-year revenue in Years 1–6; terminal NWC is zero after liquidation.","The Year 1 loss immediately reduces tax paid on other group income.","Initial NWC remains EUR 80,000 as either sensitivity input changes.","Sales scale within existing capacity; no additional CAPEX is modelled.","Past R&D and allocated existing corporate costs are excluded.","Q2 preset applies 1.10 times Q1 units and a net annual fixed-cost change of EUR -20,000; other inputs are shared."]},indent=2)+"\n")
reference = {"source":source,"wacc":ws["C38"].value,"npv":ws["C59"].value,"irr":ws["C61"].value,"depreciation":ws["C18"].value,"cashflows":[ws[f"{col}56"].value for col in "CDEFGHIJ"],"present_values":[ws[f"{col}57"].value for col in "CDEFGHIJ"]}
reference['q2']={'sheet':'Question 2','npv':q2['C61'].value,'irr':q2['C63'].value,'wacc':q2['C39'].value,'cashflows':[q2[f'{col}58'].value for col in 'CDEFGHIJ'],'present_values':[q2[f'{col}59'].value for col in 'CDEFGHIJ']}
(data/"workbook_reference.json").write_text(json.dumps(reference,indent=2)+"\n")
with (data/"baseline_cashflows.csv").open("w",newline="") as f:
    writer=csv.writer(f)
    writer.writerow(["Year","Baseline units (cans)","FCF (EUR)","PV (EUR)"])
    for i,col in enumerate("CDEFGHIJ"):
        writer.writerow([i,ws[f"{col}43"].value or 0,ws[f"{col}56"].value,ws[f"{col}57"].value])
print(json.dumps({"source":SOURCE.name,"wacc":reference["wacc"],"npv":reference["npv"],"irr":reference["irr"]}))

